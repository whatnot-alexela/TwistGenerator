"""Tests for the owner's export and stats.

The export is the research output the whole rating mechanism exists for, so
what matters is that nothing is quietly lost: not a comment with a control
character in it, not a generation nobody rated, not the cost of a call that
failed.
"""

from __future__ import annotations

from collections.abc import AsyncIterator
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from core.export import (
    ExportFilter,
    build_workbook,
    collect_stats,
    parse_window,
)
from core.quota import local_today
from db.models import ApiCall, BudgetLedger, Generation, Payment, Rating, User
from db.session import create_all, make_engine, make_session_factory

TODAY = local_today()


@pytest_asyncio.fixture
async def session() -> AsyncIterator[AsyncSession]:
    engine = make_engine("sqlite+aiosqlite://")
    await create_all(engine)
    factory = make_session_factory(engine)
    async with factory() as active:
        yield active
    await engine.dispose()


async def make_user(session: AsyncSession, telegram_id: int = 100, **kwargs: Any) -> User:
    user = User(telegram_id=telegram_id, username=f"user{telegram_id}", **kwargs)
    session.add(user)
    await session.flush()
    return user


async def make_generation(
    session: AsyncSession, user: User, formula: str = "1е-ФД", **kwargs: Any
) -> Generation:
    defaults: dict[str, Any] = {
        "mode": "formula",
        "formula": formula,
        "change_type": 1,
        "change_kind": "е",
        "cause_1": "Ф",
        "cause_2": "Д",
        "paradox_1": False,
        "paradox_2": True,
        "catalogued": True,
        "had_reference": True,
        "example_seed": 1,
        "block_b": "срез",
        "profile": "full",
        "response_text": "Ожидание. Откровение.",
        "model": "claude-opus-5",
        "effort": "high",
        "units_charged": 1,
        "was_free": True,
    }
    generation = Generation(user_id=user.id, **{**defaults, **kwargs})
    session.add(generation)
    await session.flush()
    return generation


async def make_call(
    session: AsyncSession, user: User, generation: Generation | None, cost: str, **kwargs: Any
) -> ApiCall:
    call = ApiCall(
        generation_id=generation.id if generation else None,
        user_id=user.id,
        purpose="generation",
        model="claude-opus-5",
        effort="high",
        input_tokens=kwargs.pop("input_tokens", 20_000),
        output_tokens=kwargs.pop("output_tokens", 4_000),
        cost_usd=Decimal(cost),
        was_free=True,
        **kwargs,
    )
    session.add(call)
    await session.flush()
    return call


def sheets(data: bytes) -> dict[str, list[list[Any]]]:
    book = load_workbook(BytesIO(data))
    return {
        name: [list(row) for row in book[name].iter_rows(values_only=True)]
        for name in book.sheetnames
    }


# --------------------------------------------------------------------------- #
# Stats
# --------------------------------------------------------------------------- #


async def test_stats_on_an_empty_database(session: AsyncSession) -> None:
    figures = await collect_stats(session, TODAY)
    assert figures.generations == 0
    assert figures.average_cost is None  # not a division by zero
    assert figures.average_score is None


async def test_stats_report_the_measured_average_cost(session: AsyncSession) -> None:
    """This is the number that replaces the $0.25 estimate the prices rest on."""
    user = await make_user(session)
    for _ in range(4):
        generation = await make_generation(session, user)
        await make_call(session, user, generation, "0.20")
    await session.flush()

    figures = await collect_stats(session, TODAY)
    assert figures.generations == 4
    assert figures.spend_total == Decimal("0.80")
    assert figures.average_cost == Decimal("0.20")


async def test_stats_count_failures_separately(session: AsyncSession) -> None:
    user = await make_user(session)
    generation = await make_generation(session, user)
    await make_call(session, user, generation, "0.20")
    await make_call(session, user, None, "0.08", error="refused: declined")
    await session.flush()

    figures = await collect_stats(session, TODAY)
    assert figures.failures == 1
    # The failed call still cost money, so it is still in the total.
    assert figures.spend_total == Decimal("0.28")


async def test_stats_average_the_ratings(session: AsyncSession) -> None:
    user = await make_user(session)
    for score in (5, 4, 3):
        generation = await make_generation(session, user)
        session.add(Rating(generation_id=generation.id, score=score))
    await session.flush()

    figures = await collect_stats(session, TODAY)
    assert figures.rated == 3
    assert figures.average_score == Decimal("4.00")


async def test_stats_split_today_from_the_month(session: AsyncSession) -> None:
    session.add(BudgetLedger(period_date=TODAY, free_spend_usd=Decimal("1.50")))
    if TODAY.day > 1:
        session.add(BudgetLedger(period_date=TODAY.replace(day=1), free_spend_usd=Decimal("2.00")))
    await session.flush()

    figures = await collect_stats(session, TODAY)
    assert figures.spend_today == Decimal("1.50")
    assert figures.spend_month >= Decimal("1.50")


# --------------------------------------------------------------------------- #
# The workbook
# --------------------------------------------------------------------------- #


async def test_the_workbook_has_every_sheet(session: AsyncSession) -> None:
    user = await make_user(session)
    generation = await make_generation(session, user)
    await make_call(session, user, generation, "0.20")
    await session.flush()

    report = await build_workbook(session, TODAY)
    assert set(sheets(report.data)) == {
        "generations",
        "ratings",
        "costs",
        "api_calls",
        "users",
    }
    assert report.filename.endswith(".xlsx")


async def test_a_generation_carries_its_rating_and_its_cost(session: AsyncSession) -> None:
    user = await make_user(session)
    generation = await make_generation(session, user, formula="6и-ФК")
    await make_call(session, user, generation, "0.2345")
    session.add(Rating(generation_id=generation.id, score=5, comment="отлично"))
    await session.flush()

    rows = sheets((await build_workbook(session, TODAY)).data)["generations"]
    header, row = rows[0], rows[1]
    assert row[header.index("формула")] == "6и-ФК"
    assert row[header.index("оценка")] == 5
    assert row[header.index("комментарий")] == "отлично"
    assert row[header.index("стоимость $")] == pytest.approx(0.2345)
    assert row[header.index("текст твиста")] == "Ожидание. Откровение."


async def test_an_unrated_generation_is_still_exported(session: AsyncSession) -> None:
    """Most generations will never be rated; losing them would lose the data."""
    user = await make_user(session)
    await make_generation(session, user)
    await session.flush()

    rows = sheets((await build_workbook(session, TODAY)).data)["generations"]
    assert len(rows) == 2  # header plus one
    assert rows[1][rows[0].index("оценка")] is None


async def test_the_paradox_column_is_readable(session: AsyncSession) -> None:
    user = await make_user(session)
    await make_generation(session, user, paradox_1=True, paradox_2=True)
    await session.flush()

    rows = sheets((await build_workbook(session, TODAY)).data)["generations"]
    assert rows[1][rows[0].index("парадокс")] == "оба"


async def test_ratings_are_broken_down_by_formula_and_paradox(session: AsyncSession) -> None:
    user = await make_user(session)
    first = await make_generation(session, user, formula="1е-ФД", paradox_2=True)
    second = await make_generation(session, user, formula="1е-ФМ", paradox_2=False)
    session.add(Rating(generation_id=first.id, score=5))
    session.add(Rating(generation_id=second.id, score=2))
    await session.flush()

    rows = sheets((await build_workbook(session, TODAY)).data)["ratings"]
    cuts = {row[0] for row in rows[1:]}
    assert "формула" in cuts
    assert "парадокс №2" in cuts
    assert "есть пример" in cuts


async def test_failed_calls_appear_with_their_error(session: AsyncSession) -> None:
    user = await make_user(session)
    generation = await make_generation(session, user)
    await make_call(session, user, generation, "0.20")
    await make_call(session, user, None, "0.08", error="refused: declined")
    await session.flush()

    rows = sheets((await build_workbook(session, TODAY)).data)["api_calls"]
    errors = [row[rows[0].index("ошибка")] for row in rows[1:]]
    assert "refused: declined" in errors


async def test_purchases_are_totalled_per_user(session: AsyncSession) -> None:
    user = await make_user(session)
    session.add(
        Payment(
            user_id=user.id,
            pack_id="small",
            stars=260,
            units=10,
            telegram_payment_charge_id="charge-1",
        )
    )
    await make_generation(session, user)
    await session.flush()

    rows = sheets((await build_workbook(session, TODAY)).data)["users"]
    header, row = rows[0], rows[1]
    assert row[header.index("куплено ⭐")] == 260
    assert row[header.index("генераций")] == 1


async def test_a_user_with_no_generations_still_appears(session: AsyncSession) -> None:
    await make_user(session, telegram_id=777)
    await session.flush()

    rows = sheets((await build_workbook(session, TODAY)).data)["users"]
    assert rows[1][rows[0].index("генераций")] == 0


async def test_an_empty_database_produces_an_empty_report(session: AsyncSession) -> None:
    report = await build_workbook(session, TODAY)
    assert report.rows["generations"] == 0


# --------------------------------------------------------------------------- #
# Text that Excel would reject
# --------------------------------------------------------------------------- #


async def test_a_control_character_does_not_break_the_file(session: AsyncSession) -> None:
    """A user can paste anything into a comment; openpyxl refuses some of it."""
    user = await make_user(session)
    generation = await make_generation(session, user)
    session.add(Rating(generation_id=generation.id, score=4, comment="плохо\x07\x00ок"))
    await session.flush()

    rows = sheets((await build_workbook(session, TODAY)).data)["generations"]
    comment = rows[1][rows[0].index("комментарий")]
    assert comment == "плохоок"


async def test_a_very_long_twist_is_truncated_visibly(session: AsyncSession) -> None:
    user = await make_user(session)
    await make_generation(session, user, response_text="я" * 40_000)
    await session.flush()

    rows = sheets((await build_workbook(session, TODAY)).data)["generations"]
    text = rows[1][rows[0].index("текст твиста")]
    assert len(text) < 40_000
    assert text.endswith("[обрезано]")


async def test_timestamps_survive_the_trip(session: AsyncSession) -> None:
    user = await make_user(session)
    await make_generation(session, user)
    await session.flush()

    rows = sheets((await build_workbook(session, TODAY)).data)["generations"]
    assert isinstance(rows[1][rows[0].index("время")], datetime)


# --------------------------------------------------------------------------- #
# Date windows
# --------------------------------------------------------------------------- #


def test_parse_window_accepts_none_one_or_two_dates() -> None:
    assert parse_window([]) == ExportFilter()
    assert parse_window(["2026-09-01"]).since == date(2026, 9, 1)

    window = parse_window(["2026-09-01", "2026-09-30"])
    assert window.since == date(2026, 9, 1)
    assert window.until == date(2026, 9, 30)


def test_parse_window_rejects_nonsense() -> None:
    with pytest.raises(ValueError, match="не похоже на дату"):
        parse_window(["вчера"])
    with pytest.raises(ValueError, match="позже конечной"):
        parse_window(["2026-09-30", "2026-09-01"])


async def test_a_window_filters_the_ledger(session: AsyncSession) -> None:
    session.add(BudgetLedger(period_date=TODAY, free_spend_usd=Decimal("1")))
    session.add(BudgetLedger(period_date=TODAY - timedelta(days=40), free_spend_usd=Decimal("9")))
    await session.flush()

    full = await build_workbook(session, TODAY)
    narrowed = await build_workbook(session, TODAY, ExportFilter(since=TODAY))

    assert full.rows["costs"] == 2
    assert narrowed.rows["costs"] == 1


async def test_a_window_shows_in_the_filename(session: AsyncSession) -> None:
    report = await build_workbook(session, TODAY, ExportFilter(since=date(2026, 9, 1)))
    assert "2026-09-01" in report.filename
